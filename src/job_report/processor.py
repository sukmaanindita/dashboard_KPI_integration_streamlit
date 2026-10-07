from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import BinaryIO

import pandas as pd
from openpyxl import load_workbook


STATUS_COLUMNS = {
    "success": "success",
    "retry success rate < 3 jam": "retry_lt_3h",
    "retry success rate <= 3 jam": "retry_lt_3h",
    "retry success rate > 3 jam": "retry_gt_3h",
}
PRIORITY_NAMES = {"H": "HIGH", "M": "MEDIUM", "L": "LOW"}


@dataclass(frozen=True)
class JobReport:
    jobs: pd.DataFrame
    runs: pd.DataFrame
    summary: pd.DataFrame
    validation: pd.DataFrame

    def pivot(self, period: str | None = None) -> pd.DataFrame:
        data = self.summary
        if period:
            data = data[data["period"].eq(period.upper())]
        group_cols = ["group_job", "priority", "priority_name"]
        pivot = (
            data.groupby(group_cols, dropna=False, as_index=False)
            .agg(
                job_count=("job_name", "count"),
                periodic_job=("periodic_job", "sum"),
                success=("success", "sum"),
                retry_lt_3h=("retry_lt_3h", "sum"),
                retry_gt_3h=("retry_gt_3h", "sum"),
                success_pct=("success_pct", "mean"),
                retry_lt_3h_pct=("retry_lt_3h_pct", "mean"),
                retry_gt_3h_pct=("retry_gt_3h_pct", "mean"),
                completed_pct=("completed_pct", "mean"),
            )
        )
        percentage_columns = [
            "success_pct",
            "retry_lt_3h_pct",
            "retry_gt_3h_pct",
            "completed_pct",
        ]
        pivot[percentage_columns] = pivot[percentage_columns].round(2)
        return pivot


def _clean(value: object) -> str:
    return " ".join(str(value or "").strip().split())


def _canonical_source(value: object) -> str:
    source = _clean(value).upper()
    if source == "SSIS":
        return "SSIS"
    return source


def canonical_group(value: object) -> str:
    """Standardize group names, including replacing ampersands with DAN."""
    return _clean(_clean(value).upper().replace("&", " DAN "))


def _status_key(value: object) -> str | None:
    return STATUS_COLUMNS.get(_clean(value).lower())


def add_percentages(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    denominator = result["periodic_job"].replace(0, pd.NA)
    for count_col, pct_col in (
        ("success", "success_pct"),
        ("retry_lt_3h", "retry_lt_3h_pct"),
        ("retry_gt_3h", "retry_gt_3h_pct"),
    ):
        result[pct_col] = (result[count_col] / denominator * 100).fillna(0).round(2)
    result["completed_pct"] = (
        (result["success"] + result["retry_lt_3h"] + result["retry_gt_3h"])
        / denominator
        * 100
    ).fillna(0).round(2)
    return result


def load_job_report(
    source: str | Path | BinaryIO | BytesIO,
    sheet_name: str = "Result Detail JOB",
) -> JobReport:
    workbook = load_workbook(source, data_only=True, read_only=False)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' tidak ditemukan.")
    sheet = workbook[sheet_name]

    jobs: list[dict[str, object]] = []
    runs: list[dict[str, object]] = []
    errors: list[dict[str, object]] = []

    day_columns: list[tuple[int, int, str]] = []
    for column in range(6, min(sheet.max_column, 98) + 1):
        day_value = sheet.cell(1, column).value
        if day_value is None:
            # The day header is merged across three status columns.
            day_value = sheet.cell(1, column - ((column - 6) % 3)).value
        status = _status_key(sheet.cell(2, column).value)
        if day_value and status:
            day_columns.append((column, int(day_value), status))

    if len(day_columns) != 93:
        raise ValueError(
            f"Diharapkan 31 hari x 3 status, tetapi ditemukan {len(day_columns)} kolom."
        )

    for row in range(3, sheet.max_row + 1):
        job_name = _clean(sheet.cell(row, 1).value)
        if not job_name:
            continue
        job = {
            "job_name": job_name,
            "source": _canonical_source(sheet.cell(row, 2).value),
            "period": _clean(sheet.cell(row, 3).value).upper(),
            "group_job": canonical_group(sheet.cell(row, 4).value),
            "priority": _clean(sheet.cell(row, 5).value).upper(),
        }
        job["priority_name"] = PRIORITY_NAMES.get(job["priority"], job["priority"])
        jobs.append(job)

        per_day: dict[int, list[str]] = {}
        for column, day, status in day_columns:
            marker = sheet.cell(row, column).value
            if marker in (None, "", 0, False):
                continue
            try:
                quantity = float(marker)
            except (TypeError, ValueError):
                errors.append(
                    {**job, "day": day, "issue": f"Marker tidak numerik: {marker!r}"}
                )
                continue
            if quantity <= 0:
                continue
            per_day.setdefault(day, []).append(status)
            runs.append({**job, "day": day, "status": status, "quantity": quantity})

        for day, statuses in per_day.items():
            if len(statuses) > 1:
                errors.append(
                    {
                        **job,
                        "day": day,
                        "issue": "Lebih dari satu status: " + ", ".join(statuses),
                    }
                )
        if job["period"] not in {"DAILY", "MONTHLY"}:
            errors.append({**job, "day": None, "issue": "Period tidak dikenal"})
        if job["priority"] not in PRIORITY_NAMES:
            errors.append({**job, "day": None, "issue": "Priority tidak dikenal"})

    jobs_df = pd.DataFrame(jobs).drop_duplicates().reset_index(drop=True)
    runs_df = pd.DataFrame(runs)
    validation_df = pd.DataFrame(errors)

    keys = ["job_name", "source", "period", "group_job", "priority", "priority_name"]
    if runs_df.empty:
        counts = pd.DataFrame(columns=keys + ["periodic_job", "success", "retry_lt_3h", "retry_gt_3h"])
    else:
        counts = (
            runs_df.pivot_table(
                index=keys,
                columns="status",
                values="quantity",
                aggfunc="sum",
                fill_value=0,
            )
            .reset_index()
        )
        for column in ("success", "retry_lt_3h", "retry_gt_3h"):
            if column not in counts:
                counts[column] = 0
        counts["periodic_job"] = counts[["success", "retry_lt_3h", "retry_gt_3h"]].sum(axis=1)

    summary = jobs_df.merge(counts, on=keys, how="left")
    for column in ("periodic_job", "success", "retry_lt_3h", "retry_gt_3h"):
        summary[column] = summary[column].fillna(0)
    summary = add_percentages(summary)
    return JobReport(jobs_df, runs_df, summary, validation_df)
